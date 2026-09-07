"""Offline o'lchangan hujjat/olchov Excel va statistikada SYNC vaqti emas,
HAQIQIY o'lchangan kuni bo'yicha ko'rinishi kerak.

- HujjatCreate/OlchovCreate: `olchandi_vaqt` maydoni
- _xavfsiz_olchov_vaqti(): ishonchsiz (kelajak / 30 kundan eski) qiymatni rad
  etib datetime.now() qaytaradi
- POST /hujjatlar: created_at VA hujjat raqamining yili `olchandi_vaqt` dan
- POST /olchovlar: yangi qatorda created_at `olchandi_vaqt` dan; MAVJUD
  qatorni yangilashda created_at o'zgarmaydi
"""
import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

import main
from models import Hujjat, HujjatHolati


# ---------- _xavfsiz_olchov_vaqti ----------

def test_none_hozirni_qaytaradi():
    oldin = datetime.now()
    r = main._xavfsiz_olchov_vaqti(None)
    assert oldin <= r <= datetime.now() + timedelta(seconds=1)


def test_kelajakdagi_vaqt_rad_etiladi():
    kelajak = datetime.now() + timedelta(days=2)
    r = main._xavfsiz_olchov_vaqti(kelajak)
    assert r < kelajak  # hozirga qaytarildi


def test_30_kundan_eski_vaqt_rad_etiladi():
    juda_eski = datetime.now() - timedelta(days=45)
    r = main._xavfsiz_olchov_vaqti(juda_eski)
    assert r > juda_eski  # hozirga qaytarildi


def test_yaroqli_otgan_vaqt_ozgarishsiz_qaytadi():
    kecha = datetime.now() - timedelta(days=1, hours=3)
    assert main._xavfsiz_olchov_vaqti(kecha) == kecha


def test_timezone_aware_vaqt_naive_ga_keltiriladi():
    aware = datetime.now(timezone.utc) - timedelta(hours=2)
    r = main._xavfsiz_olchov_vaqti(aware)
    assert r.tzinfo is None
    assert r < datetime.now()


# ---------- POST /hujjatlar ----------

def test_hujjat_created_at_olchandi_vaqtdan(client, admin_headers, mahsulot_chigit, mashina):
    olchov_vaqti = datetime.now() - timedelta(days=3, hours=5)
    javob = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
        "olchandi_vaqt": olchov_vaqti.isoformat(),
    }, headers=admin_headers)
    assert javob.status_code == 200
    got = datetime.fromisoformat(javob.json()["created_at"])
    assert abs((got - olchov_vaqti).total_seconds()) < 2


def test_hujjat_raqami_yili_olchandi_vaqtdan(
        client, admin_headers, mahsulot_chigit, mashina, monkeypatch):
    # 31-Dekabrda o'lchangan, 2-Yanvarda sync bo'lgan holatni simulyatsiya
    # qilamiz (30 kunlik chegara ichida): raqam O'TGAN yilni olishi kerak.
    otgan_yil_dek = datetime(datetime.now().year - 1, 12, 31, 12, 0)
    monkeypatch.setattr("main._xavfsiz_olchov_vaqti", lambda v: otgan_yil_dek)

    javob = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
    }, headers=admin_headers)
    assert javob.status_code == 200
    assert f"-{otgan_yil_dek.year}/" in javob.json()["raqam"]      # CHG-2025/001
    assert javob.json()["created_at"].startswith(str(otgan_yil_dek.year))


def test_olchandi_vaqtsiz_created_at_hozir(client, admin_headers, mahsulot_chigit, mashina):
    oldin = datetime.now() - timedelta(seconds=2)
    javob = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
    }, headers=admin_headers)
    got = datetime.fromisoformat(javob.json()["created_at"])
    assert oldin <= got <= datetime.now() + timedelta(seconds=2)


# ---------- POST /olchovlar ----------

def test_olchov_yangi_qator_created_at_olchandi_vaqtdan(
        client, admin_headers, mahsulot_chigit, mashina):
    h = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
    }, headers=admin_headers).json()
    olchov_vaqti = datetime.now() - timedelta(days=2)
    client.post("/olchovlar", json={
        "hujjat_id": h["id"], "arava_raqam": 1, "tara": 1000,
        "olchandi_vaqt": olchov_vaqti.isoformat(),
    }, headers=admin_headers)
    olchovlar = client.get(f"/olchovlar/{h['id']}", headers=admin_headers).json()
    got = datetime.fromisoformat(olchovlar[0]["created_at"])
    assert abs((got - olchov_vaqti).total_seconds()) < 2


def test_olchov_yangilash_created_at_ni_ozgartirmaydi(
        client, admin_headers, mahsulot_chigit, mashina):
    h = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
    }, headers=admin_headers).json()
    tara_vaqti = datetime.now() - timedelta(days=2)
    # 1) tara - yangi qator, created_at = tara_vaqti
    client.post("/olchovlar", json={
        "hujjat_id": h["id"], "arava_raqam": 1, "tara": 1000,
        "olchandi_vaqt": tara_vaqti.isoformat(),
    }, headers=admin_headers)
    # 2) brutto - MAVJUD qator yangilanadi, olchandi_vaqt=hozir bo'lsa ham
    #    created_at o'zgarmasligi kerak (birinchi o'lchov vaqtini bildiradi)
    client.post("/olchovlar", json={
        "hujjat_id": h["id"], "arava_raqam": 1, "brutto": 5000,
        "olchandi_vaqt": datetime.now().isoformat(),
    }, headers=admin_headers)
    olchovlar = client.get(f"/olchovlar/{h['id']}", headers=admin_headers).json()
    assert len(olchovlar) == 1
    got = datetime.fromisoformat(olchovlar[0]["created_at"])
    assert abs((got - tara_vaqti).total_seconds()) < 2
    assert olchovlar[0]["brutto"] == 5000


# ---------- Integratsiya: Excel va statistika ----------

def test_offline_hujjat_eksportda_olchangan_kuni_bilan(
        client, admin_headers, mahsulot_chigit, mashina):
    uch_kun_oldin = datetime.now() - timedelta(days=3)
    h = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
        "olchandi_vaqt": uch_kun_oldin.isoformat(),
    }, headers=admin_headers).json()
    client.post("/olchovlar", json={
        "hujjat_id": h["id"], "arava_raqam": 1, "tara": 1000, "brutto": 5000,
    }, headers=admin_headers)

    kun = uch_kun_oldin.date()
    er = client.get(
        f"/hujjatlar/eksport?mahsulot_id={mahsulot_chigit.id}"
        f"&sana_dan={kun}&sana_gacha={kun + timedelta(days=1)}",
        headers=admin_headers)
    assert er.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(er.content))
    matnlar = [str(c.value) for row in wb.active.iter_rows() for c in row]
    assert h["raqam"] in matnlar, "offline hujjat o'lchangan kunidagi eksportda yo'q"

    # Bugungi eksportda esa BO'LMASLIGI kerak
    bugun = datetime.now().date()
    er2 = client.get(
        f"/hujjatlar/eksport?mahsulot_id={mahsulot_chigit.id}"
        f"&sana_dan={bugun}&sana_gacha={bugun + timedelta(days=1)}",
        headers=admin_headers)
    matnlar2 = [str(c.value) for row in openpyxl.load_workbook(
        io.BytesIO(er2.content)).active.iter_rows() for c in row]
    assert h["raqam"] not in matnlar2


def test_offline_hujjat_kunlik_statistikaga_kirmaydi(
        client, admin_headers, db_session, mahsulot_chigit, mashina):
    # 3 kun oldin o'lchangan, tugallangan hujjat -> bugungi kunlik
    # statistikaga kirmasligi kerak (created_at bugungi emas)
    uch_kun_oldin = datetime.now() - timedelta(days=3)
    h = client.post("/hujjatlar", json={
        "mashina_id": mashina.id, "mahsulot_id": mahsulot_chigit.id,
        "olchandi_vaqt": uch_kun_oldin.isoformat(),
    }, headers=admin_headers).json()
    client.post("/olchovlar", json={
        "hujjat_id": h["id"], "arava_raqam": 1, "tara": 1000, "brutto": 5000,
    }, headers=admin_headers)
    # holatni tugallandi qilamiz (statistika filtri: TUGALLANDI + netto>0)
    hujjat = db_session.query(Hujjat).filter(Hujjat.id == h["id"]).first()
    hujjat.holat = HujjatHolati.TUGALLANDI
    db_session.commit()

    kunlik = client.get("/statistika/kunlik", headers=admin_headers).json()
    assert kunlik["jami_tonnaj"] == 0.0   # bugun o'lchangan narsa yo'q

    haftalik = client.get("/statistika/haftalik", headers=admin_headers).json()
    assert haftalik["jami_tonnaj"] == pytest.approx(4.0)  # 5000-1000 = 4 t
